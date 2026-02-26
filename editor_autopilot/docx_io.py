from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass
class ParsedDoc:
    document: Document
    paragraphs: list[str]
    tables: list[list[list[str]]]


def load_docx(path: str) -> ParsedDoc:
    p = Path(path)
    if p.suffix.lower() != ".docx":
        raise ValueError("Only .docx is supported. Scanned PDF is not supported.")
    doc = Document(path)
    paragraphs = [para.text.strip() for para in doc.paragraphs]
    tables: list[list[list[str]]] = []
    for tbl in doc.tables:
        rows: list[list[str]] = []
        for row in tbl.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return ParsedDoc(document=doc, paragraphs=paragraphs, tables=tables)


def save_docx(doc: Document, path: str) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
